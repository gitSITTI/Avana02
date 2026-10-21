"""CSV / TSV / XLSX / HTML tables: one requirement per row. Also used by DOORS 'export to CSV/Excel' and any ALM export.

HTML: a Digital.ai Agility grid "Export (.xls)" is NOT a workbook. The vendor's documentation says it is an HTML table saved with an .xls
extension, and the MSOSA / Cameo requirement table can be exported as .html. Both are read here: every top-level <table> is one table
(named after its <caption>, else `<file>#<n>`; choose one with the `sheet` option). A cell becomes its visible text (tags dropped, <br> <p>
<li> become line breaks, entities decoded); a table nested in a cell becomes text of that cell; colspan keeps the columns aligned; rowspan is
NOT expanded. A real binary .xls (the old Excel format) is not read.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterator

from ..adapter import Adapter, Caps, Container, Discovery, Source, make_field_info, register
from ..model import Fidelity, RawRecord, SourceRef


def file_sha(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


def find_header(rows: list[list[str]], scan: int = 15) -> int:
    """Header = first row (within the first `scan`) that is mostly filled, mostly non-numeric, and followed by data.
    Skips title blocks and blank lines that real spreadsheets have above the table."""
    for i, row in enumerate(rows[:scan]):
        cells = [c.strip() for c in row]
        filled = [c for c in cells if c]
        if len(filled) < 2 or len(filled) < 0.6 * len(cells):
            continue
        numeric = sum(1 for c in filled if c.replace(".", "", 1).isdigit())
        if numeric > 0.3 * len(filled):
            continue
        if i + 1 < len(rows) and any(c.strip() for c in rows[i + 1]):
            return i
    return 0


def _read_csv(path: str) -> list[list[str]]:
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(path: str) -> dict[str, list[list[str]]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)  # data_only: cached values, never formulas
    out: dict[str, list[list[str]]] = {}
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        out[ws.title] = [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return out


_HTML_EXT = (".html", ".htm", ".xls")
_BLOCK_TAGS = {"br", "p", "div", "li", "ul", "ol", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "pre", "blockquote"}


def _html_head(path: str, n: int = 65536) -> str:
    try:
        with open(path, "rb") as f:
            return f.read(n).decode("utf-8", errors="replace").lower()
    except OSError:
        return ""


def looks_like_html_table(path: str) -> bool:
    """True when the start of the file is HTML (not a binary Excel file, not a zip) and the first 64 KB hold a <table."""
    try:
        with open(path, "rb") as f:
            raw = f.read(65536)
    except OSError:
        return False
    if raw.startswith((b"\xd0\xcf\x11\xe0", b"PK")):          # binary .xls (OLE2) or a zip: not HTML
        return False
    return b"<table" in raw.lower()


def _decode_html(raw: bytes) -> str:
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8", errors="replace")
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")
    m = re.search(rb"<meta[^>]+charset\s*=\s*[\"']?\s*([A-Za-z0-9_.:-]+)", raw[:4096], re.I)
    for enc in ([m.group(1).decode("ascii", "replace")] if m else []) + ["utf-8", "cp1252"]:
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("latin-1")


def _clean_cell(parts: list[str]) -> str:
    text = "".join(parts).replace("\xa0", " ").replace("\u200b", "")
    lines = [re.sub(r"[ \t\r\f\v]+", " ", ln).strip() for ln in text.split("\n")]
    return "\n".join(ln for ln in lines if ln)


class _TableParser(HTMLParser):
    """Collects the rows of every top-level <table>. Tolerant of omitted </td> and </tr>, which HTML allows and exporters use."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[tuple[str, list[list[str]]]] = []
        self._depth = 0
        self._rows: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._span = 1
        self._caption: list[str] | None = None
        self._caption_text = ""
        self._skip = 0

    def _close_cell(self) -> None:
        if self._cell is not None and self._row is not None:
            self._row.append(_clean_cell(self._cell))
            self._row.extend([""] * (self._span - 1))
        self._cell, self._span = None, 1

    def _close_row(self) -> None:
        self._close_cell()
        if self._row is not None and self._rows is not None:
            self._rows.append(self._row)
        self._row = None

    def handle_starttag(self, tag: str, attrs) -> None:  # noqa: ANN001
        if tag in ("script", "style"):
            self._skip += 1
            return
        if tag == "table":
            self._depth += 1
            if self._depth == 1:
                self._rows, self._caption_text = [], ""
            elif self._cell is not None:
                self._cell.append("\n")
            return
        if self._depth == 1:
            if tag == "caption":
                self._caption = []
            elif tag == "tr":
                self._close_row()
                self._row = []
            elif tag in ("td", "th"):
                if self._row is None:
                    self._row = []
                self._close_cell()
                try:
                    self._span = max(1, min(50, int(dict(attrs).get("colspan") or 1)))
                except ValueError:
                    self._span = 1
                self._cell = []
        if self._cell is not None and tag in _BLOCK_TAGS:
            self._cell.append("\n")

    def handle_startendtag(self, tag: str, attrs) -> None:  # noqa: ANN001
        if tag in _BLOCK_TAGS and self._cell is not None:
            self._cell.append("\n")
        elif tag in ("td", "th", "tr", "table"):
            self.handle_starttag(tag, attrs)
            self.handle_endtag(tag)

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style"):
            self._skip = max(0, self._skip - 1)
            return
        if tag == "caption" and self._caption is not None:
            self._caption_text, self._caption = _clean_cell(self._caption).replace("\n", " "), None
        elif tag in ("td", "th") and self._depth == 1:
            self._close_cell()
        elif tag == "tr" and self._depth == 1:
            self._close_row()
        elif tag == "table" and self._depth > 0:
            if self._depth == 1:
                self._close_row()
                self.tables.append((self._caption_text, self._rows or []))
                self._rows = None
            self._depth -= 1
        if self._cell is not None and tag in _BLOCK_TAGS:
            self._cell.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip:
            return
        if self._caption is not None:
            self._caption.append(data)
        elif self._cell is not None:
            self._cell.append(data)

    def close(self) -> None:
        super().close()
        if self._depth:                         # a file cut off in the middle of a table keeps what it had
            self._close_row()
            if self._rows is not None:
                self.tables.append((self._caption_text, self._rows))


def _read_html(path: str) -> dict[str, list[list[str]]]:
    p = _TableParser()
    p.feed(_decode_html(Path(path).read_bytes()))
    p.close()
    out: dict[str, list[list[str]]] = {}
    for n, (caption, rows) in enumerate(p.tables, start=1):
        name = caption or f"{Path(path).stem}#{n}"
        if name in out:
            name = f"{name}#{n}"
        out[name] = rows
    return out


@register
class TabularAdapter(Adapter):
    name = "tabular"
    fidelity = Fidelity.TABULAR
    caps = Caps()

    def sniff(self, src: Source) -> float:
        ext = Path(src.uri).suffix.lower()
        if ext in (".csv", ".tsv"):
            return 0.9
        if ext in (".xlsx", ".xlsm"):
            try:
                return 0.95 if "xl/workbook.xml" in zipfile.ZipFile(src.uri).namelist() else 0.0
            except (zipfile.BadZipFile, OSError):
                return 0.0
        if ext in _HTML_EXT:                      # .xls is claimed only when it is really an HTML table, as an Agility grid export is
            return 0.9 if looks_like_html_table(src.uri) else 0.0
        return 0.0

    def _tables(self, src: Source) -> dict[str, list[dict[str, str]]]:
        ext = Path(src.uri).suffix.lower()
        if ext in (".csv", ".tsv"):
            sheets = {Path(src.uri).stem: _read_csv(src.uri)}
        elif ext in _HTML_EXT:
            sheets = _read_html(src.uri)
        else:
            sheets = _read_xlsx(src.uri)
        want = src.options.get("sheet")
        if want and ext in _HTML_EXT and want not in sheets and str(want).isdigit() and 1 <= int(want) <= len(sheets):
            want = list(sheets)[int(want) - 1]          # a table can also be chosen by its position: sheet=2
        out: dict[str, list[dict[str, str]]] = {}
        for name, rows in sheets.items():
            if want and name != want:
                continue
            rows = [r for r in rows]
            if not any(any(c.strip() for c in r) for r in rows):
                continue
            h = src.options.get("header_row")
            h = find_header(rows) if h is None else int(h)
            header = [c.strip() or f"col{j + 1}" for j, c in enumerate(rows[h])]
            seen: dict[str, int] = {}
            for j, c in enumerate(header):           # make duplicate headers unique instead of losing a column
                seen[c] = seen.get(c, 0) + 1
                if seen[c] > 1:
                    header[j] = f"{c}#{seen[c]}"
            body = []
            for k, r in enumerate(rows[h + 1:], start=h + 2):  # 1-based sheet row numbers
                if not any(c.strip() for c in r):
                    continue
                d = {header[j]: (r[j] if j < len(r) else "") for j in range(len(header))}
                d["_row"] = str(k)
                body.append(d)
            out[name] = body
        return out

    def discover(self, src: Source) -> Discovery:
        cs = []
        for name, rows in self._tables(src).items():
            cols = [k for k in rows[0] if k != "_row"] if rows else []
            cs.append(Container(name, len(rows), [make_field_info(c, [r.get(c, "") for r in rows]) for c in cols]))
        return Discovery(cs, ["header row auto-detected; override with options['header_row'] (0-based)"])

    def read(self, src: Source) -> Iterator[RawRecord]:
        rev = self.revision(src)
        for name, rows in self._tables(src).items():
            for i, r in enumerate(rows):
                fields = {k: v for k, v in r.items() if k != "_row"}
                yield RawRecord(native_id=None, fields=fields, order=i,
                                ref=SourceRef("tabular", src.uri, container=name, locator=f"{name}!R{r['_row']}", revision=rev))

    def revision(self, src: Source) -> str:
        return file_sha(src.uri)

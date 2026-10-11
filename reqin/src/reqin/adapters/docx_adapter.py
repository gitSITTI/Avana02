"""Word (.docx): tables become tabular containers (best case: an SRS table); prose paragraphs become EXTRACTED candidates.

A .doc (old binary) must be converted first (LibreOffice: soffice --headless --convert-to docx).
"""
from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Iterator

from ..adapter import Adapter, Caps, Container, Discovery, Source, make_field_info, register
from ..model import Fidelity, RawRecord, SourceRef
from .prose import HEADING, candidates
from .tabular import file_sha, find_header


@register
class DocxAdapter(Adapter):
    name = "docx"
    fidelity = Fidelity.EXTRACTED
    caps = Caps()

    def sniff(self, src: Source) -> float:
        if Path(src.uri).suffix.lower() != ".docx":
            return 0.0
        try:
            return 0.95 if "word/document.xml" in zipfile.ZipFile(src.uri).namelist() else 0.0
        except (zipfile.BadZipFile, OSError):
            return 0.0

    def _walk(self, src: Source):
        import docx
        d = docx.Document(src.uri)
        section = ""
        prose: list[dict] = []
        for i, p in enumerate(d.paragraphs):
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            if style.lower().startswith("heading") or HEADING.match(t):
                section = t
                continue
            cands = candidates(t)
            for k, (lead, stmt) in enumerate(cands, start=1):
                loc = f"p{i + 1}" + (f".{k}" if len(cands) > 1 else "")   # unique locator per statement
                prose.append({"text": stmt, "lead_id": lead or "", "section": section, "_loc": loc})
        tables: dict[str, list[dict]] = {}
        for ti, tb in enumerate(d.tables, start=1):
            rows = [[c.text.strip() for c in r.cells] for r in tb.rows]
            if len(rows) < 2:
                continue
            h = find_header(rows)
            header = [c or f"col{j + 1}" for j, c in enumerate(rows[h])]
            body = []
            for k, r in enumerate(rows[h + 1:], start=h + 2):
                if any(r):
                    body.append({**{header[j]: r[j] for j in range(min(len(header), len(r)))}, "_loc": f"t{ti}!R{k}"})
            tables[f"table{ti}"] = body
        return prose, tables

    def discover(self, src: Source) -> Discovery:
        prose, tables = self._walk(src)
        cs = []
        if prose:
            cs.append(Container("prose", len(prose), [make_field_info(c, [r[c] for r in prose]) for c in ("text", "lead_id", "section")]))
        for name, rows in tables.items():
            cols = [k for k in rows[0] if k != "_loc"]
            cs.append(Container(name, len(rows), [make_field_info(c, [r.get(c, "") for r in rows]) for c in cols]))
        return Discovery(cs, ["prose candidates are sentences/paragraphs containing shall/must/should/will; verify each against the document"])

    def read(self, src: Source) -> Iterator[RawRecord]:
        prose, tables = self._walk(src)
        rev = self.revision(src)
        n = 0
        for r in prose:
            yield RawRecord(native_id=None, fields={k: v for k, v in r.items() if k != "_loc"}, order=n,
                            ref=SourceRef("docx", src.uri, "prose", r["_loc"], rev))
            n += 1
        for name, rows in tables.items():
            for r in rows:
                yield RawRecord(native_id=None, fields={k: v for k, v in r.items() if k != "_loc"}, order=n,
                                ref=SourceRef("docx", src.uri, name, r["_loc"], rev))
                n += 1

    def revision(self, src: Source) -> str:
        return file_sha(src.uri)

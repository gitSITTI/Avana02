"""Builds one synthetic requirement set in every supported format. Fictional data, no program content.

The set is designed so that:
  * REQ-FCS-0101 / 0102 are the same obligation in different units (200 ms vs 0.2 seconds)
  * REQ-FCS-0103 / 0104 differ in modal (shall vs should) and value (30 vs 14 days)
  * REQ-MC-0301 / 0302 are paraphrases
so a downstream comparison step can be shown to work ACROSS formats.
REAL exports differ. Before trusting an adapter on your data, add an anonymised real export as a fixture.
"""
from __future__ import annotations

import csv
import shutil
import sqlite3
import subprocess
from pathlib import Path

REQS = [
    # uid, text, rationale, verification, allocation, owner, parent
    ("REQ-FCS-0101", "The Flight Control Subsystem shall transmit a status message to the ground station within 200 ms of a mode change.", "Bounds ground-station loop latency", "Test", "Flight Control Subsystem", "FCS team", ""),
    ("REQ-FCS-0102", "The Flight Control Subsystem shall transmit a status message to the ground station within 0.2 seconds of a mode change.", "", "Test", "Flight Control Subsystem", "FCS team", "REQ-FCS-0101"),
    ("REQ-FCS-0103", "The Flight Control Subsystem shall retain at least 30 days of fault records.", "Maintenance planning", "Inspection", "Flight Control Subsystem", "FCS team", ""),
    ("REQ-FCS-0104", "The Flight Control Subsystem should retain at least 14 days of fault records.", "", "Inspection", "Flight Control Subsystem", "FCS team", "REQ-FCS-0103"),
    ("REQ-PWR-0201", "The Power Subsystem shall not exceed a mass of 12 kg.", "Mass budget", "Analysis", "Power Subsystem", "PWR team", ""),
    ("REQ-PWR-0202", "The Power Subsystem shall provide 28 VDC to the Mission Computer.", "", "Test", "Power Subsystem", "PWR team", ""),
    ("REQ-MC-0301", "The Mission Computer shall log each operator command with a timestamp.", "", "Demonstration", "Mission Computer", "MC team", ""),
    ("REQ-MC-0302", "The Mission Computer shall log every operator command together with the time it was received.", "", "Demonstration", "Mission Computer", "MC team", ""),
]
COUNT = len(REQS)


def build(out: Path) -> dict[str, Path]:
    out.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    # CSV: clean headers
    p = out / "reqs.csv"
    with p.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Requirement ID", "Statement", "Rationale", "Verification Method", "Allocated To", "Owner", "Parent ID"])
        for r in REQS:
            w.writerow([r[0], r[1], r[2], r[3], r[4], r[5], r[6]])
    paths["csv"] = p

    # XLSX: a Cover sheet (not requirements), three title rows above the header, a blank row, a hidden sheet
    import openpyxl
    wb = openpyxl.Workbook()
    cover = wb.active; cover.title = "Cover"
    cover["A1"] = "Fictional program"; cover["A2"] = "Owner"; cover["B2"] = "Nobody"
    ws = wb.create_sheet("Requirements")
    ws["A1"] = "Subsystem Requirements Spec (fictional)"; ws["A2"] = "Draft"; ws["A3"] = ""
    ws.append(["Req #", "Requirement Text", "Why", "V Method", "Subsystem", "IPT"])
    for r in REQS[:5]:
        ws.append([r[0], r[1], r[2], {"Test": "T", "Inspection": "I", "Analysis": "A", "Demonstration": "D"}[r[3]], r[4], r[5]])
    ws.append([None] * 6)
    for r in REQS[5:]:
        ws.append([r[0], r[1], r[2], {"Test": "T", "Inspection": "I", "Analysis": "A", "Demonstration": "D"}[r[3]], r[4], r[5]])
    hid = wb.create_sheet("Scratch"); hid["A1"] = "ignore me"; hid.sheet_state = "hidden"
    p = out / "reqs.xlsx"; wb.save(p); paths["xlsx"] = p

    # DOORS-style CSV export: standard DOORS attribute names, heading rows without text
    p = out / "doors_export.csv"
    with p.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Absolute Number", "Object Identifier", "Object Level", "Object Heading", "Object Text", "Verification Method"])
        w.writerow([1, "SRS-1", 1, "Flight control", "", ""])
        for i, r in enumerate(REQS[:4], start=2):
            w.writerow([i, r[0], 2, "", r[1], r[3]])
        w.writerow([6, "SRS-2", 1, "Power", "", ""])
        for i, r in enumerate(REQS[4:6], start=7):
            w.writerow([i, r[0], 2, "", r[1], r[3]])
    paths["doors_csv"] = p

    paths["reqif"] = _reqif(out / "reqs.reqif")
    paths["sysml"] = _sysml(out / "model.sysml")
    paths["xmi"] = _xmi(out / "model.xmi")
    paths["mdzip"] = _mdzip(out / "model.mdzip", paths["xmi"])

    # SQLite
    p = out / "reqs.db"
    p.unlink(missing_ok=True)
    con = sqlite3.connect(p)
    con.execute("CREATE TABLE reqs (id INTEGER PRIMARY KEY, req_uid TEXT, statement TEXT, rationale TEXT, verif TEXT, alloc TEXT, owner TEXT, active INTEGER)")
    for i, r in enumerate(REQS, start=1):
        con.execute("INSERT INTO reqs VALUES (?,?,?,?,?,?,?,1)", (i, r[0], r[1], r[2], r[3], r[4], r[5]))
    con.execute("INSERT INTO reqs VALUES (99,'REQ-OLD-0001','The legacy subsystem shall be retired.','','Test','Legacy','X',0)")
    con.commit(); con.close()
    paths["sqlite"] = p

    paths["docx"] = _docx(out / "srs.docx")
    paths["pdf"] = _pdf(out / "srs.pdf")               # pure Python (reportlab): same content as the docx, no LibreOffice needed
    if shutil.which("soffice"):                         # OPTIONAL extra: the docx converted by a real word processor
        office = out / "office"
        office.mkdir(exist_ok=True)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(office), str(paths["docx"])],
                       check=True, capture_output=True, timeout=180)
        if (office / "srs.pdf").exists():
            paths["pdf_office"] = office / "srs.pdf"
    return paths


def _sysml(p: Path) -> Path:
    def d(i):  # doc comment, wrapped like a human would
        t = REQS[i][1]
        k = len(t) // 2
        k = t.rfind(" ", 0, k)
        return f"/* {t[:k]}\n           {t[k + 1:]} */"
    p.write_text(f"""// Fictional SysML v2 textual model. The word requirement appears in this comment on purpose.
package FixtureModel {{
    requirement def <'REQ-FCS-0101'> StatusLatency {{ doc {d(0)} }}
    requirement def <'REQ-FCS-0102'> StatusLatencySeconds {{ doc {d(1)} }}
    requirement def <'REQ-FCS-0103'> FaultRetention {{
        doc {d(2)}
        requirement def <'REQ-FCS-0104'> FaultRetentionShould {{ doc {d(3)} }}
    }}
    requirement def <'REQ-PWR-0201'> PowerMass {{ doc {d(4)} }}
    requirement def <'REQ-PWR-0202'> PowerSupply {{ doc {d(5)} }}
    requirement def <'REQ-MC-0301'> CommandLog {{ doc {d(6)} }}
    requirement def <'REQ-MC-0302'> CommandLogTime {{ doc {d(7)} }}
    part def FlightControl;
    part flightControl : FlightControl {{
        satisfy requirement StatusLatency by flightControl;
    }}
    requirement latencyUsage : StatusLatency;
}}
""", encoding="utf-8")
    return p


def _xmi(p: Path) -> Path:
    """A fictional SysML 1.x model as XMI 2.4, shaped as the vendor's documentation describes (never checked against a real file): requirement classes in a
    package, nested where REQS says there is a parent, a sibling Requirement application per class, one Block, one Satisfy and one DeriveReqt."""
    from xml.sax.saxutils import quoteattr as q
    ids = {r[0]: f"_c{n}" for n, r in enumerate(REQS, start=1)}
    def cls(r, kids: str) -> str:
        return f'<{"nestedClassifier" if r[6] else "packagedElement"} xmi:type="uml:Class" xmi:id="{ids[r[0]]}" name={q(r[0] + " requirement")}>{kids}</{"nestedClassifier" if r[6] else "packagedElement"}>'
    def tree(r) -> str:
        return cls(r, "".join(tree(k) for k in REQS if k[6] == r[0]))
    body = "".join(tree(r) for r in REQS if not r[6])
    apps = "".join(
        f'<sysml:Requirement xmi:id="_s{n}" base_Class="{ids[r[0]]}" Id={q(r[0])} Text={q(r[1])} Rationale={q(r[2])} VerifyMethod={q(r[3])}/>'
        for n, r in enumerate(REQS, start=1))
    p.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<xmi:XMI xmi:version="2.4" xmlns:xmi="http://www.omg.org/spec/XMI/20110701" xmlns:uml="http://www.omg.org/spec/UML/20110701" '
        'xmlns:sysml="http://www.omg.org/spec/SysML/20110701/SysML">'
        f'<uml:Model xmi:id="_m" name="Fixture Model"><packagedElement xmi:type="uml:Package" xmi:id="_p" name="Subsystem requirements">{body}'
        '<packagedElement xmi:type="uml:Class" xmi:id="_b1" name="Flight Control block"/>'
        f'<packagedElement xmi:type="uml:Abstraction" xmi:id="_a1" client="_b1" supplier="{ids[REQS[0][0]]}"/>'
        f'<packagedElement xmi:type="uml:Abstraction" xmi:id="_a2" client="{ids[REQS[1][0]]}" supplier="{ids[REQS[0][0]]}"/>'
        '</packagedElement></uml:Model>'
        f'{apps}<sysml:Block xmi:id="_sb" base_Class="_b1"/><sysml:Satisfy xmi:id="_sa1" base_Abstraction="_a1"/><sysml:DeriveReqt xmi:id="_sa2" base_Abstraction="_a2"/>'
        '</xmi:XMI>', encoding="utf-8")
    return p


def _mdzip(p: Path, xmi: Path) -> Path:
    """The layout the vendor documents from 17.0.1: the model in one entry, other resources beside it."""
    import zipfile
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("com.nomagic.magicdraw.uml_model.model", xmi.read_bytes())
        z.writestr("com.nomagic.magicdraw.uml_model.shared_model", "<xmi:XMI xmlns:xmi='http://www.omg.org/spec/XMI/20110701'/>")
        z.writestr("images/fixture.png", b"\x89PNG fixture")
    return p


def _reqif(p: Path) -> Path:
    T = "2026-01-01T00:00:00Z"
    enum_vals = {"Test": "ev-t", "Analysis": "ev-a", "Inspection": "ev-i", "Demonstration": "ev-d"}
    ev = "".join(f'<ENUM-VALUE IDENTIFIER="{i}" LAST-CHANGE="{T}" LONG-NAME="{n}"><PROPERTIES><EMBEDDED-VALUE KEY="{k}" OTHER-CONTENT=""/></PROPERTIES></ENUM-VALUE>'
                 for k, (n, i) in enumerate(enum_vals.items()))
    def sstr(defn, val):
        return f'<ATTRIBUTE-VALUE-STRING THE-VALUE="{val}"><DEFINITION><ATTRIBUTE-DEFINITION-STRING-REF>{defn}</ATTRIBUTE-DEFINITION-STRING-REF></DEFINITION></ATTRIBUTE-VALUE-STRING>'
    objs = []
    objs.append(f'<SPEC-OBJECT IDENTIFIER="so-h1" LAST-CHANGE="{T}"><TYPE><SPEC-OBJECT-TYPE-REF>sot-head</SPEC-OBJECT-TYPE-REF></TYPE><VALUES>{sstr("ad-chap", "Subsystem requirements")}</VALUES></SPEC-OBJECT>')
    for n, r in enumerate(REQS, start=1):
        xhtml = f'<xhtml:div><xhtml:p>{r[1]}</xhtml:p></xhtml:div>' if n % 2 else f'<xhtml:div>{r[1]}</xhtml:div>'
        objs.append(
            f'<SPEC-OBJECT IDENTIFIER="so-{n}" LAST-CHANGE="{T}"><TYPE><SPEC-OBJECT-TYPE-REF>sot-req</SPEC-OBJECT-TYPE-REF></TYPE><VALUES>'
            f'{sstr("ad-id", r[0])}'
            f'<ATTRIBUTE-VALUE-XHTML><DEFINITION><ATTRIBUTE-DEFINITION-XHTML-REF>ad-text</ATTRIBUTE-DEFINITION-XHTML-REF></DEFINITION><THE-VALUE>{xhtml}</THE-VALUE></ATTRIBUTE-VALUE-XHTML>'
            f'{sstr("ad-why", r[2])}'
            f'<ATTRIBUTE-VALUE-ENUMERATION><DEFINITION><ATTRIBUTE-DEFINITION-ENUMERATION-REF>ad-ver</ATTRIBUTE-DEFINITION-ENUMERATION-REF></DEFINITION><VALUES><ENUM-VALUE-REF>{enum_vals[r[3]]}</ENUM-VALUE-REF></VALUES></ATTRIBUTE-VALUE-ENUMERATION>'
            f'</VALUES></SPEC-OBJECT>')
    kids = "".join(f'<SPEC-HIERARCHY IDENTIFIER="h-{n}" LAST-CHANGE="{T}"><OBJECT><SPEC-OBJECT-REF>so-{n}</SPEC-OBJECT-REF></OBJECT></SPEC-HIERARCHY>' for n in range(1, COUNT + 1))
    xml = f'''<?xml version="1.0" encoding="UTF-8"?>
<REQ-IF xmlns="http://www.omg.org/spec/ReqIF/20110401/reqif.xsd" xmlns:xhtml="http://www.w3.org/1999/xhtml">
<THE-HEADER><REQ-IF-HEADER IDENTIFIER="hdr"><CREATION-TIME>{T}</CREATION-TIME><REQ-IF-TOOL-ID>fixture</REQ-IF-TOOL-ID><REQ-IF-VERSION>1.0</REQ-IF-VERSION><SOURCE-TOOL-ID>fixture</SOURCE-TOOL-ID><TITLE>fixture</TITLE></REQ-IF-HEADER></THE-HEADER>
<CORE-CONTENT><REQ-IF-CONTENT>
<DATATYPES>
<DATATYPE-DEFINITION-STRING IDENTIFIER="dt-str" LAST-CHANGE="{T}" LONG-NAME="String" MAX-LENGTH="32000"/>
<DATATYPE-DEFINITION-XHTML IDENTIFIER="dt-xhtml" LAST-CHANGE="{T}" LONG-NAME="XHTML"/>
<DATATYPE-DEFINITION-ENUMERATION IDENTIFIER="dt-ver" LAST-CHANGE="{T}" LONG-NAME="Verification"><SPECIFIED-VALUES>{ev}</SPECIFIED-VALUES></DATATYPE-DEFINITION-ENUMERATION>
</DATATYPES>
<SPEC-TYPES>
<SPEC-OBJECT-TYPE IDENTIFIER="sot-req" LAST-CHANGE="{T}" LONG-NAME="Requirement"><SPEC-ATTRIBUTES>
<ATTRIBUTE-DEFINITION-STRING IDENTIFIER="ad-id" LAST-CHANGE="{T}" LONG-NAME="ReqIF.ForeignID"><TYPE><DATATYPE-DEFINITION-STRING-REF>dt-str</DATATYPE-DEFINITION-STRING-REF></TYPE></ATTRIBUTE-DEFINITION-STRING>
<ATTRIBUTE-DEFINITION-XHTML IDENTIFIER="ad-text" LAST-CHANGE="{T}" LONG-NAME="ReqIF.Text"><TYPE><DATATYPE-DEFINITION-XHTML-REF>dt-xhtml</DATATYPE-DEFINITION-XHTML-REF></TYPE></ATTRIBUTE-DEFINITION-XHTML>
<ATTRIBUTE-DEFINITION-STRING IDENTIFIER="ad-why" LAST-CHANGE="{T}" LONG-NAME="Rationale"><TYPE><DATATYPE-DEFINITION-STRING-REF>dt-str</DATATYPE-DEFINITION-STRING-REF></TYPE></ATTRIBUTE-DEFINITION-STRING>
<ATTRIBUTE-DEFINITION-ENUMERATION IDENTIFIER="ad-ver" LAST-CHANGE="{T}" LONG-NAME="Verification Method" MULTI-VALUED="false"><TYPE><DATATYPE-DEFINITION-ENUMERATION-REF>dt-ver</DATATYPE-DEFINITION-ENUMERATION-REF></TYPE></ATTRIBUTE-DEFINITION-ENUMERATION>
</SPEC-ATTRIBUTES></SPEC-OBJECT-TYPE>
<SPEC-OBJECT-TYPE IDENTIFIER="sot-head" LAST-CHANGE="{T}" LONG-NAME="Heading"><SPEC-ATTRIBUTES>
<ATTRIBUTE-DEFINITION-STRING IDENTIFIER="ad-chap" LAST-CHANGE="{T}" LONG-NAME="ReqIF.ChapterName"><TYPE><DATATYPE-DEFINITION-STRING-REF>dt-str</DATATYPE-DEFINITION-STRING-REF></TYPE></ATTRIBUTE-DEFINITION-STRING>
</SPEC-ATTRIBUTES></SPEC-OBJECT-TYPE>
<SPEC-RELATION-TYPE IDENTIFIER="srt-derive" LAST-CHANGE="{T}" LONG-NAME="Derive"/>
<SPECIFICATION-TYPE IDENTIFIER="st-1" LAST-CHANGE="{T}" LONG-NAME="Spec"/>
</SPEC-TYPES>
<SPEC-OBJECTS>{"".join(objs)}</SPEC-OBJECTS>
<SPEC-RELATIONS><SPEC-RELATION IDENTIFIER="rel-1" LAST-CHANGE="{T}"><TYPE><SPEC-RELATION-TYPE-REF>srt-derive</SPEC-RELATION-TYPE-REF></TYPE><SOURCE><SPEC-OBJECT-REF>so-2</SPEC-OBJECT-REF></SOURCE><TARGET><SPEC-OBJECT-REF>so-1</SPEC-OBJECT-REF></TARGET></SPEC-RELATION></SPEC-RELATIONS>
<SPECIFICATIONS><SPECIFICATION IDENTIFIER="spec-1" LAST-CHANGE="{T}" LONG-NAME="Subsystem requirements"><TYPE><SPECIFICATION-TYPE-REF>st-1</SPECIFICATION-TYPE-REF></TYPE><CHILDREN>
<SPEC-HIERARCHY IDENTIFIER="h-0" LAST-CHANGE="{T}"><OBJECT><SPEC-OBJECT-REF>so-h1</SPEC-OBJECT-REF></OBJECT><CHILDREN>{kids}</CHILDREN></SPEC-HIERARCHY>
</CHILDREN></SPECIFICATION></SPECIFICATIONS>
</REQ-IF-CONTENT></CORE-CONTENT></REQ-IF>
'''
    p.write_text(xml, encoding="utf-8")
    return p


def _docx(p: Path) -> Path:
    import docx
    d = docx.Document()
    d.add_heading("3 Performance requirements", level=1)
    d.add_paragraph("This section describes the performance expectations in plain language and contains no obligations.")
    for r in REQS[:2]:
        d.add_paragraph(f"{r[0]}: {r[1]}")
    d.add_paragraph(f"{REQS[2][0]}: {REQS[2][1]} {REQS[3][1]}")      # two obligations in one paragraph (no IDs on the second)
    d.add_heading("4 Power and computing", level=1)
    t = d.add_table(rows=1, cols=3); t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["ID", "Requirement", "Verification"]):
        c.text = h
    for r in REQS[4:]:
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = r[0], r[1], r[3]
    d.save(p)
    return p


def _pdf(p: Path) -> Path:
    """Same content as _docx: prose with IDs, a paragraph carrying two obligations, and a ruled table."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    st = getSampleStyleSheet()
    body, h1 = st["BodyText"], st["Heading1"]
    cell = st["BodyText"].clone("cell", fontSize=9, leading=11)
    esc = lambda t: t.replace("&", "&amp;").replace("<", "&lt;")  # noqa: E731
    doc = SimpleDocTemplate(str(p), pagesize=letter, invariant=1, title="Fictional SRS", author="fixture")
    flow = [Paragraph("3 Performance requirements", h1),
            Paragraph("This section describes the performance expectations in plain language and contains no obligations.", body)]
    for r in REQS[:2]:
        flow.append(Paragraph(esc(f"{r[0]}: {r[1]}"), body))
    flow.append(Paragraph(esc(f"{REQS[2][0]}: {REQS[2][1]} {REQS[3][1]}"), body))
    flow += [Spacer(1, 12), Paragraph("4 Power and computing", h1)]
    rows = [["ID", "Requirement", "Verification"]] + [[Paragraph(esc(r[0]), cell), Paragraph(esc(r[1]), cell), Paragraph(r[3], cell)] for r in REQS[4:]]
    t = Table(rows, colWidths=[90, 330, 80])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.8, colors.black), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]))
    flow.append(t)
    doc.build(flow)
    return p
